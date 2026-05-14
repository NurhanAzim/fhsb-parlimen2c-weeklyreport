from __future__ import annotations

import json
import posixpath
import uuid
from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from docx.image.image import Image as DocxImage


ODT_NS = {
    "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
    "manifest": "urn:oasis:names:tc:opendocument:xmlns:manifest:1.0",
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "style": "urn:oasis:names:tc:opendocument:xmlns:style:1.0",
    "svg": "urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "xlink": "http://www.w3.org/1999/xlink",
}

for prefix, uri in ODT_NS.items():
    ET.register_namespace(prefix, uri)


WORK_IMAGE_MAX_WIDTH_IN = 3.4375
WORK_IMAGE_MAX_HEIGHT_IN = 4.75
FIRST_PAGE_WORK_TOKEN = "<work_detail>"
FIRST_PAGE_IMAGE_TOKEN = "<work_images>"
FIRST_PAGE_NOTE_TOKEN = "<work_note>"
LAST_PAGE_WORK_TOKEN = "<work2_detail>"
LAST_PAGE_IMAGE_TOKEN = "<work2_images>"
LAST_PAGE_NOTE_TOKEN = "<work2_note>"


@dataclass(slots=True)
class Issue:
    description: str
    images_description: str = ""
    image_paths: list[Path] = field(default_factory=list)


@dataclass(slots=True)
class ReportData:
    date_start: str
    date_end: str
    location: str
    issues: list[Issue] = field(default_factory=list)

    @classmethod
    def from_dict(cls, payload: dict) -> "ReportData":
        issues = [
            Issue(
                description=item["description"].strip(),
                images_description=item.get("images_description", "").strip(),
                image_paths=[Path(path) for path in item.get("image_paths", [])],
            )
            for item in payload.get("issues", [])
        ]
        return cls(
            date_start=payload["date_start"].strip(),
            date_end=payload["date_end"].strip(),
            location=payload["location"].strip(),
            issues=issues,
        )

    def placeholder_map(self) -> dict[str, str]:
        return {
            "date_start": self.date_start,
            "date_end": self.date_end,
            "month_year": month_year_from_date(self.date_start),
            "location": self.location,
        }


def month_year_from_date(value: str) -> str:
    month_map = {
        "01": "JANUARI",
        "02": "FEBRUARI",
        "03": "MAC",
        "04": "APRIL",
        "05": "MEI",
        "06": "JUN",
        "07": "JULAI",
        "08": "OGOS",
        "09": "SEPTEMBER",
        "10": "OKTOBER",
        "11": "NOVEMBER",
        "12": "DISEMBER",
    }
    _day, month, year = value.strip().split("/")
    return f"{month_map[month]} {year}"


def render_report(
    template_path: str | Path,
    output_path: str | Path,
    report: ReportData,
) -> Path:
    template = Path(template_path)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    with ZipFile(template) as source:
        package = {name: source.read(name) for name in source.namelist()}

    content_root = ET.fromstring(package["content.xml"])
    styles_root = ET.fromstring(package["styles.xml"])
    manifest_root = ET.fromstring(package["META-INF/manifest.xml"])
    _replace_scalar_placeholders(styles_root, report.placeholder_map())
    _populate_content(content_root, styles_root, report)

    picture_entries = _collect_picture_entries(report.issues or [Issue(description="", images_description="", image_paths=[])])
    _update_manifest(manifest_root, picture_entries)

    package["content.xml"] = ET.tostring(content_root, encoding="utf-8", xml_declaration=True)
    package["styles.xml"] = ET.tostring(styles_root, encoding="utf-8", xml_declaration=True)
    package["META-INF/manifest.xml"] = ET.tostring(manifest_root, encoding="utf-8", xml_declaration=True)
    for package_path, source_path in picture_entries:
        package[package_path] = source_path.read_bytes()

    _write_odt(output, package, package.keys())
    return output


def _replace_scalar_placeholders(root: ET.Element, replacements: dict[str, str]) -> None:
    xml_text = ET.tostring(root, encoding="unicode")
    for name, value in replacements.items():
        xml_text = xml_text.replace(f"&lt;{name}&gt;", _xml_escape(value))
    replacement_root = ET.fromstring(xml_text)
    root.clear()
    root.tag = replacement_root.tag
    root.attrib.update(replacement_root.attrib)
    for child in list(replacement_root):
        root.append(child)


def _populate_content(root: ET.Element, styles_root: ET.Element, report: ReportData) -> None:
    _replace_scalar_placeholders(root, report.placeholder_map())
    entries = _page_entries_for_issues(
        report.issues or [Issue(description="", images_description="", image_paths=[])]
    )

    table = root.find(".//table:table", ODT_NS)
    if table is None:
        raise ValueError("Weekly report table not found in ODT template.")
    table_parent = _find_parent(root, table)
    if table_parent is None:
        raise ValueError("Weekly report table parent not found in ODT template.")

    body_rows = list(table.findall("table:table-row", ODT_NS))
    if len(body_rows) < 2:
        raise ValueError("Unexpected ODT row layout for weekly report template.")

    first_row_template = deepcopy(body_rows[0])
    last_row_template = deepcopy(body_rows[1])
    body_table_style_name, last_table_style_name = _ensure_page_styles(root, styles_root, table)
    insert_at = list(table_parent).index(table)
    table_parent.remove(table)

    for index, entry in enumerate(entries):
        is_last_page = index == len(entries) - 1
        page_table = deepcopy(table)
        _clear_table_body(page_table)
        page_table.attrib[_q("table", "style-name")] = last_table_style_name if is_last_page else body_table_style_name

        row = deepcopy(last_row_template if (len(entries) == 1 or is_last_page) else first_row_template)
        _fill_work_row(row, entry, use_second_page_tokens=(len(entries) == 1 or is_last_page))
        page_table.append(row)

        table_parent.insert(insert_at + index, page_table)


def _page_entries_for_issues(issues: list[Issue]) -> list[Issue]:
    entries: list[Issue] = []
    for issue in issues:
        if not issue.image_paths:
            entries.append(
                Issue(
                    description=issue.description,
                    images_description=issue.images_description,
                    image_paths=[],
                )
            )
            continue

        for image_path in issue.image_paths:
            entries.append(
                Issue(
                    description=issue.description,
                    images_description=issue.images_description,
                    image_paths=[image_path],
                )
            )
    return entries


def _ensure_page_styles(
    content_root: ET.Element,
    styles_root: ET.Element,
    table: ET.Element,
) -> tuple[str, str]:
    content_styles = content_root.find("office:automatic-styles", ODT_NS)
    if content_styles is None:
        raise ValueError("ODT automatic styles not found in content.xml.")

    master_styles = styles_root.find("office:master-styles", ODT_NS)
    if master_styles is None:
        raise ValueError("ODT master styles not found in styles.xml.")

    base_table_style_name = table.attrib.get(_q("table", "style-name"))
    if not base_table_style_name:
        raise ValueError("Weekly report table style missing in ODT template.")

    base_table_style = content_styles.find(f"style:style[@style:name='{base_table_style_name}']", ODT_NS)
    if base_table_style is None:
        raise ValueError("Weekly report table style definition not found in ODT template.")

    master_page = master_styles.find("style:master-page", ODT_NS)
    if master_page is None:
        raise ValueError("ODT master page not found in styles.xml.")

    body_master_name = "ReportBody"
    last_master_name = "ReportLast"
    body_table_style_name = f"{base_table_style_name}.Body"
    last_table_style_name = f"{base_table_style_name}.Last"

    body_master = _clone_master_page(master_page, body_master_name)
    _set_master_page_footer_mode(body_master, blank_footer=True)
    _upsert_named_child(master_styles, body_master, "style", "master-page", body_master_name)

    last_master = _clone_master_page(master_page, last_master_name)
    _set_master_page_footer_mode(last_master, blank_footer=False)
    _upsert_named_child(master_styles, last_master, "style", "master-page", last_master_name)
    _set_master_page_footer_mode(master_page, blank_footer=True)

    body_table_style = deepcopy(base_table_style)
    body_table_style.attrib[_q("style", "name")] = body_table_style_name
    body_table_style.attrib[_q("style", "master-page-name")] = body_master_name
    _upsert_named_child(content_styles, body_table_style, "style", "style", body_table_style_name)

    last_table_style = deepcopy(base_table_style)
    last_table_style.attrib[_q("style", "name")] = last_table_style_name
    last_table_style.attrib[_q("style", "master-page-name")] = last_master_name
    _upsert_named_child(content_styles, last_table_style, "style", "style", last_table_style_name)

    return body_table_style_name, last_table_style_name


def _clone_master_page(master_page: ET.Element, style_name: str) -> ET.Element:
    clone = deepcopy(master_page)
    clone.attrib[_q("style", "name")] = style_name
    return clone


def _set_master_page_footer_mode(master_page: ET.Element, blank_footer: bool) -> None:
    header = master_page.find("style:header", ODT_NS)
    header_first = master_page.find("style:header-first", ODT_NS)
    if header is not None:
        if header_first is None:
            header_first = ET.SubElement(master_page, _q("style", "header-first"))
        _replace_children(header_first, header)

    footer = master_page.find("style:footer", ODT_NS)
    footer_first = master_page.find("style:footer-first", ODT_NS)
    if footer is None and not blank_footer:
        return

    if footer is None:
        footer = ET.SubElement(master_page, _q("style", "footer"))
    if footer_first is None:
        footer_first = ET.SubElement(master_page, _q("style", "footer-first"))

    if blank_footer:
        _clear_node_content(footer)
        _clear_node_content(footer_first)
        ET.SubElement(footer, _q("text", "p"), {_q("text", "style-name"): "Footer"})
        ET.SubElement(footer_first, _q("text", "p"), {_q("text", "style-name"): "Footer"})
        return

    _replace_children(footer_first, footer)


def _clear_table_body(table: ET.Element) -> None:
    for child in list(table):
        if child.tag in {_q("table", "table-row"), _q("text", "soft-page-break")}:
            table.remove(child)


def _fill_work_row(row: ET.Element, entry: Issue, use_second_page_tokens: bool) -> None:
    cells = row.findall("table:table-cell", ODT_NS)
    if len(cells) != 3:
        raise ValueError("Unexpected ODT work row shape.")

    detail_token = LAST_PAGE_WORK_TOKEN if use_second_page_tokens else FIRST_PAGE_WORK_TOKEN
    image_token = LAST_PAGE_IMAGE_TOKEN if use_second_page_tokens else FIRST_PAGE_IMAGE_TOKEN
    note_token = LAST_PAGE_NOTE_TOKEN if use_second_page_tokens else FIRST_PAGE_NOTE_TOKEN

    _set_cell_text(cells[0], detail_token, _normalize_work_detail(entry.description))
    _set_cell_text(cells[2], note_token, _normalize_work_note(entry.images_description))
    _set_image_cell(cells[1], image_token, entry.image_paths)


def _set_cell_text(cell: ET.Element, token: str, value: str) -> None:
    paragraphs = cell.findall("text:p", ODT_NS)
    target = paragraphs[-1] if paragraphs else ET.SubElement(cell, _q("text", "p"))
    for paragraph in paragraphs[:-1]:
        _clear_node_content(paragraph)
    _clear_node_content(target)
    target.text = value


def _set_image_cell(cell: ET.Element, token: str, image_paths: list[Path]) -> None:
    paragraphs = cell.findall("text:p", ODT_NS)
    template = paragraphs[-1] if paragraphs else ET.SubElement(cell, _q("text", "p"))
    for paragraph in paragraphs:
        cell.remove(paragraph)

    if not image_paths:
        paragraph = deepcopy(template)
        _clear_node_content(paragraph)
        cell.append(paragraph)
        return

    for image_path in image_paths:
        paragraph = deepcopy(template)
        _clear_node_content(paragraph)
        paragraph.append(_build_image_frame(image_path))
        cell.append(paragraph)


def _build_image_frame(image_path: Path) -> ET.Element:
    width_in, height_in = _scaled_dimensions_in_inches(image_path)
    frame = ET.Element(
        _q("draw", "frame"),
        {
            _q("draw", "name"): image_path.stem,
            _q("text", "anchor-type"): "as-char",
            _q("svg", "width"): f"{width_in:.4f}in",
            _q("svg", "height"): f"{height_in:.4f}in",
            _q("draw", "z-index"): "0",
        },
    )
    ET.SubElement(
        frame,
        _q("draw", "image"),
        {
            _q("xlink", "href"): _package_picture_path(image_path),
            _q("xlink", "type"): "simple",
            _q("xlink", "show"): "embed",
            _q("xlink", "actuate"): "onLoad",
        },
    )
    return frame


def _collect_picture_entries(entries: list[Issue]) -> list[tuple[str, Path]]:
    seen: dict[str, Path] = {}
    for entry in entries:
        for image_path in entry.image_paths:
            seen[_package_picture_path(image_path)] = image_path
    return sorted(seen.items())


def _update_manifest(manifest_root: ET.Element, picture_entries: list[tuple[str, Path]]) -> None:
    existing = {
        element.attrib.get(_q("manifest", "full-path")): element
        for element in manifest_root.findall("manifest:file-entry", ODT_NS)
    }
    for package_path, source_path in picture_entries:
        if package_path in existing:
            existing[package_path].attrib[_q("manifest", "media-type")] = _media_type_for_path(source_path)
            continue
        ET.SubElement(
            manifest_root,
            _q("manifest", "file-entry"),
            {
                _q("manifest", "full-path"): package_path,
                _q("manifest", "media-type"): _media_type_for_path(source_path),
            },
        )


def _write_odt(output_path: Path, package: dict[str, bytes], member_names) -> None:
    ordered_names = ["mimetype"] + [name for name in member_names if name != "mimetype"]
    with ZipFile(output_path, "w") as target:
        for name in ordered_names:
            compression = ZIP_STORED if name == "mimetype" else ZIP_DEFLATED
            target.writestr(name, package[name], compress_type=compression)


def _scaled_dimensions_in_inches(image_path: Path) -> tuple[float, float]:
    image = DocxImage.from_file(str(image_path))
    native_width_in = int(image.width) / 914400
    native_height_in = int(image.height) / 914400

    if native_width_in <= WORK_IMAGE_MAX_WIDTH_IN and native_height_in <= WORK_IMAGE_MAX_HEIGHT_IN:
        return native_width_in, native_height_in

    width_scale = WORK_IMAGE_MAX_WIDTH_IN / native_width_in
    height_scale = WORK_IMAGE_MAX_HEIGHT_IN / native_height_in
    scale = min(width_scale, height_scale)
    return native_width_in * scale, native_height_in * scale


def _package_picture_path(image_path: Path) -> str:
    return posixpath.join("Pictures", image_path.name)


def _media_type_for_path(path: Path) -> str:
    suffix = path.suffix.lower()
    return {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".gif": "image/gif",
        ".bmp": "image/bmp",
        ".webp": "image/webp",
    }.get(suffix, "application/octet-stream")


def _normalize_work_detail(value: str) -> str:
    return value.strip().upper()


def _normalize_work_note(value: str) -> str:
    normalized = value.strip()
    if not normalized:
        return ""
    return normalized[:1].upper() + normalized[1:]


def _clear_node_content(node: ET.Element) -> None:
    node.text = None
    for child in list(node):
        node.remove(child)


def _replace_children(target: ET.Element, source: ET.Element) -> None:
    _clear_node_content(target)
    target.text = source.text
    target.tail = source.tail
    for child in list(source):
        target.append(deepcopy(child))


def _find_parent(root: ET.Element, node: ET.Element) -> ET.Element | None:
    for parent in root.iter():
        for child in list(parent):
            if child is node:
                return parent
    return None


def _upsert_named_child(
    parent: ET.Element,
    child: ET.Element,
    prefix: str,
    local_name: str,
    style_name: str,
) -> None:
    existing = parent.find(f"{prefix}:{local_name}[@style:name='{style_name}']", ODT_NS)
    if existing is not None:
        parent.remove(existing)
    parent.append(child)


def _xml_escape(value: str) -> str:
    return (
        value.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _q(prefix: str, name: str) -> str:
    return f"{{{ODT_NS[prefix]}}}{name}"


def load_report_data(payload_path: str | Path) -> ReportData:
    payload = json.loads(Path(payload_path).read_text(encoding="utf-8"))
    return ReportData.from_dict(payload)


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description="Render the weekly report ODT from JSON data.")
    parser.add_argument("payload", help="Path to a JSON file that matches ReportData.from_dict().")
    parser.add_argument(
        "--template",
        default="template-weekly-report.odt",
        help="Path to the ODT template.",
    )
    parser.add_argument(
        "--output",
        default="output/weekly-report.odt",
        help="Where to save the rendered ODT.",
    )
    args = parser.parse_args()

    report = load_report_data(args.payload)
    output = render_report(args.template, args.output, report)
    print(output)


if __name__ == "__main__":
    main()
